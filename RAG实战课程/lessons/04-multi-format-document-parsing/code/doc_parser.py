"""
多格式文档解析器
支持 PDF、Word、Excel、PPT、图片等格式
"""

import os
from pathlib import Path
from typing import Union, List, Dict, Optional
import json


# ========== PDF 解析 ==========
def extract_pdf_pymupdf(pdf_path: str) -> List[Dict]:
    """
    使用 PyMuPDF 提取 PDF 文本

    Args:
        pdf_path: PDF 文件路径

    Returns:
        每页文本内容的列表
    """
    import fitz

    doc = fitz.open(pdf_path)
    pages_content = []

    for page_num in range(len(doc)):
        page = doc[page_num]

        # 纯文本提取
        text = page.get_text("text")

        pages_content.append(
            {"page_num": page_num + 1, "text": text, "char_count": len(text)}
        )

    doc.close()
    return pages_content


def extract_pdf_with_tables(pdf_path: str) -> List[Dict]:
    """
    使用 Pdfplumber 提取 PDF 文本和表格
    适合财务、数据类 PDF
    """
    import pdfplumber

    all_content = []

    with pdfplumber.open(pdf_path) as pdf:
        for page_num, page in enumerate(pdf.pages):
            text = page.extract_text()
            tables = page.extract_tables()

            all_content.append(
                {
                    "page_num": page_num + 1,
                    "text": text or "",
                    "tables": tables or [],
                    "table_count": len(tables) if tables else 0,
                }
            )

    return all_content


# ========== Word 解析 ==========
def extract_word(docx_path: str) -> Dict:
    """
    使用 python-docx 提取 Word 文档

    Args:
        docx_path: Word 文件路径

    Returns:
        包含段落和表格的字典
    """
    from docx import Document

    doc = Document(docx_path)

    content = {"paragraphs": [], "tables": [], "total_chars": 0}

    # 提取段落
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content["paragraphs"].append({"text": text, "style": para.style.name})
            content["total_chars"] += len(text)

    # 提取表格
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)
        content["tables"].append(table_data)

    return content


# ========== Excel 解析 ==========
def extract_excel(excel_path: str) -> Dict:
    """
    使用 pandas 读取 Excel

    Args:
        excel_path: Excel 文件路径

    Returns:
        包含各工作表数据的字典
    """
    import pandas as pd

    content = {"sheets": {}, "total_rows": 0}

    excel_file = pd.ExcelFile(excel_path)

    for sheet_name in excel_file.sheet_names:
        df = pd.read_excel(excel_file, sheet_name=sheet_name)

        content["sheets"][sheet_name] = {
            "rows": df.shape[0],
            "columns": df.shape[1],
            "column_names": df.columns.tolist(),
            "data_preview": df.head(5).to_dict(),
        }
        content["total_rows"] += df.shape[0]

    return content


# ========== PPT 解析 ==========
def extract_ppt(ppt_path: str) -> Dict:
    """
    使用 python-pptx 提取 PPT 内容

    Args:
        ppt_path: PPT 文件路径

    Returns:
        包含各幻灯片内容的字典
    """
    from pptx import Presentation

    prs = Presentation(ppt_path)

    content = {"slides": [], "total_slides": len(prs.slides)}

    for slide_num, slide in enumerate(prs.slides):
        slide_content = {"slide_num": slide_num + 1, "title": None, "content": []}

        if slide.shapes.title:
            slide_content["title"] = slide.shapes.title.text

        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_content["content"].append(text)

        content["slides"].append(slide_content)

    return content


# ========== 图片 OCR ==========
def ocr_image(image_path: str, lang: str = "chi_sim+eng") -> str:
    """
    使用 RapidOCR 进行文字识别

    Args:
        image_path: 图片路径
        lang: 语言选项

    Returns:
        识别的文本
    """
    try:
        from rapidocr_onnxruntime import RapidOCR

        ocr_engine = RapidOCR()
        result, _ = ocr_engine(image_path)

        if result:
            return "\n".join([line[1] for line in result])
        return ""
    except ImportError:
        # 备用：使用 pytesseract
        import pytesseract
        from PIL import Image

        image = Image.open(image_path)
        return pytesseract.image_to_string(image, lang=lang)


# ========== 统一加载器 ==========
class DocumentLoader:
    """
    统一文档加载器
    自动识别文件格式并调用相应解析器
    """

    def __init__(self):
        self.loaders = {
            ".pdf": self._load_pdf,
            ".docx": self._load_word,
            ".doc": self._load_word,
            ".xlsx": self._load_excel,
            ".xls": self._load_excel,
            ".pptx": self._load_ppt,
            ".ppt": self._load_ppt,
            ".txt": self._load_txt,
            ".md": self._load_markdown,
            ".png": self._load_image,
            ".jpg": self._load_image,
            ".jpeg": self._load_image,
        }

    def load(self, file_path: Union[str, Path]) -> Dict:
        """加载文档"""
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {path}")

        suffix = path.suffix.lower()

        if suffix not in self.loaders:
            raise ValueError(f"不支持的格式: {suffix}")

        return self.loaders[suffix](path)

    def load_batch(self, directory: Union[str, Path]) -> List[Dict]:
        """批量加载目录下的所有文档"""
        path = Path(directory)
        results = []

        for file_path in path.rglob("*"):
            if file_path.is_file():
                try:
                    content = self.load(file_path)
                    results.append(content)
                    print(f"✓ {file_path.name}")
                except Exception as e:
                    print(f"✗ {file_path.name}: {e}")

        return results

    def _load_pdf(self, path: Path) -> Dict:
        pages = extract_pdf_pymupdf(str(path))
        return {
            "source": str(path),
            "type": "pdf",
            "pages": len(pages),
            "content": "\n\n".join([p["text"] for p in pages]),
            "details": pages,
        }

    def _load_word(self, path: Path) -> Dict:
        content = extract_word(str(path))
        return {
            "source": str(path),
            "type": "word",
            "paragraphs": len(content["paragraphs"]),
            "tables": len(content["tables"]),
            "content": "\n".join([p["text"] for p in content["paragraphs"]]),
            "details": content,
        }

    def _load_excel(self, path: Path) -> Dict:
        content = extract_excel(str(path))
        return {
            "source": str(path),
            "type": "excel",
            "sheets": list(content["sheets"].keys()),
            "content": str(content),
            "details": content,
        }

    def _load_ppt(self, path: Path) -> Dict:
        content = extract_ppt(str(path))
        all_text = []
        for slide in content["slides"]:
            if slide["title"]:
                all_text.append(slide["title"])
            all_text.extend(slide["content"])

        return {
            "source": str(path),
            "type": "ppt",
            "slides": content["total_slides"],
            "content": "\n".join(all_text),
            "details": content,
        }

    def _load_txt(self, path: Path) -> Dict:
        return {
            "source": str(path),
            "type": "text",
            "content": path.read_text(encoding="utf-8"),
        }

    def _load_markdown(self, path: Path) -> Dict:
        return {
            "source": str(path),
            "type": "markdown",
            "content": path.read_text(encoding="utf-8"),
        }

    def _load_image(self, path: Path) -> Dict:
        text = ocr_image(str(path))
        return {"source": str(path), "type": "image", "content": text}


# ========== 主函数 ==========
if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("用法: python doc_parser.py <文件路径或目录>")
        sys.exit(1)

    target = sys.argv[1]
    loader = DocumentLoader()

    if Path(target).is_dir():
        print(f"批量加载目录: {target}")
        results = loader.load_batch(target)
        print(f"\n共处理 {len(results)} 个文档")
    else:
        print(f"加载文件: {target}")
        result = loader.load(target)
        print(f"类型: {result['type']}")
        print(f"内容长度: {len(result['content'])} 字符")
